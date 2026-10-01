"""
Прогон по папке месяца (этап 2.2b): фото из <месяц>/фото, результат в
<месяц>/результат, показания и лог — в базе, одна транзакция на фото, в конце
выгрузка. Номер у нескольких абонентов → question/serial_ambiguous. Дата
показания — из имени файла WhatsApp.

Модели — фейки: фото — PNG-байты, пиксель [0,0] — номер записи в PHOTOS
(имя файла не важно, поэтому фото можно называть как в WhatsApp).
"""
import csv
import datetime as dt
from pathlib import Path

import cv2
import numpy as np
import openpyxl
import pytest

import gmr
import reader
from src.gmr.application import month
from src.gmr.domain import Outcome
from src.gmr.domain.photo_date import date_from_filename, reading_date
from src.gmr.ml import loader
from src.gmr.render import read_image
from src.gmr.storage.month import MonthDB, MonthFolder
from tests._fixtures import (
    FakeDigitDetector, FakeDigitOCR, FakeMeterDetector, FakeSerialOCR,
    make_digit_crops_with_centers, make_meter_crops, process_photo,
)

# номер -> (серийник, 5 цифр или None = цифры не прочитаны)
PHOTOS = {
    1: ("11111", "01200"),     # A-1
    2: ("22222", "04000"),     # A-2 (MINUS)
    3: ("33333", None),        # A-3, DIGITS_ERROR
    5: ("44444", "00500"),     # номер у A-4 и A-5
    7: ("55555", "00009"),     # A-6, показание уже было в таблице
}
TABLE = [  # номер, л/с, последние, текущие
    ("11111", "A-1", "1000", ""), ("22222", "A-2", "5000", ""), ("33333", "A-3", "100", ""),
    ("44444", "A-4", "10", ""), ("44444", "A-5", "20", ""), ("55555", "A-6", "7", "8"),
]
TODAY = dt.date.today().strftime("%d %m %Y")


class _Meter:
    def process_image(self, img_path, **kwargs):
        t = int(read_image(img_path)[0, 0, 0])
        return [
            {"class": "gas_meter", "conf": .95, "angle": 0., "bbox": (0, 0, 200, 40),
             "crop": np.full((40, 200, 3), t, np.uint8), "path": None},
            {"class": "serial_number", "conf": .95, "angle": 0., "bbox": (0, 40, 100, 60),
             "crop": np.full((20, 100, 3), t, np.uint8), "path": None},
        ]


class _Digits:
    def process_array(self, img, **kwargs):
        digits = PHOTOS[int(img[0, 0, 0])][1]
        if digits is None:
            return []
        return [{"class": "digit", "conf": .95, "angle": 0., "bbox": (i * 40, 0, i * 40 + 20, 30),
                 "crop": np.full((30, 20, 3), int(d), np.uint8), "path": None}
                for i, d in enumerate(digits)]


class _DigitOCR:
    def predict(self, img):
        return {"digit": int(img[0, 0, 0]), "confidence": .95}


class _SerialOCR:
    crash_on = None

    def predict_with_details(self, img):
        t = int(img[0, 0, 0])
        if t == self.crash_on:
            raise KeyboardInterrupt
        return {"text": PHOTOS[t][0], "avg_confidence": .95, "details": []}


@pytest.fixture
def ocr(monkeypatch):
    o = _SerialOCR()
    monkeypatch.setattr(loader, "YOLOInferer",
                        lambda m, conf_thresh=.8, straighten=True, output_dir=None:
                        _Digits() if straighten is False else _Meter())
    monkeypatch.setattr(loader, "CNNInferer", lambda *a, **k: _DigitOCR())
    monkeypatch.setattr(loader, "CRNNInferer", lambda *a, **k: o)
    return o


def _photo(path: Path, n: int):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(cv2.imencode(".png", np.full((100, 200, 3), n, np.uint8))[1].tobytes())


def _month(tmp_path) -> MonthFolder:
    table = tmp_path / "table.csv"
    with open(table, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Номер счетчика", "Лицевой счет", "Последние показания", "Текущие показания"])
        w.writerows(TABLE)
    month.load_table(str(tmp_path / "Октябрь"), str(table))
    return MonthFolder(tmp_path / "Октябрь")


def _run(folder: MonthFolder):
    reader.run_pipeline(reader.PipelineConfig(month_dir=str(folder.root)))


def _readings(folder):
    with MonthDB(folder.db) as db:
        return db.readings()


def _log_list(folder):
    with MonthDB(folder.db) as db:
        return db.log_rows()


def _log(folder):
    with MonthDB(folder.db) as db:
        return {r["original_filename"]: r for r in db.log_rows()}


# ─── Дата из имени файла ─────────────────────────────────────────────────────

@pytest.mark.parametrize("name, date", [
    ("IMG-20261001-WA0012.jpg", dt.date(2026, 10, 1)),
    ("WhatsApp Image 2026-09-30 at 14.32.05.jpeg", dt.date(2026, 9, 30)),
    ("1300022665.jpg", None), ("p1.jpg", None), ("IMG-20261399-WA0001.jpg", None),
    ("IMG-20260231-WA0001.jpg", None),                       # 31 февраля — не дата
])
def test_date_from_filename(name, date):
    assert date_from_filename(name) == date


def test_reading_date_falls_back_to_today():
    assert reading_date("p1.jpg", dt.date(2026, 10, 5)) == "05 10 2026"
    assert reading_date("IMG-20261001-WA0012.jpg", dt.date(2026, 10, 5)) == "01 10 2026"


# ─── Прогон по месяцу ────────────────────────────────────────────────────────

def test_month_run(tmp_path, ocr):
    f = _month(tmp_path)
    _photo(f.photos / "Аюб" / "IMG-20260915-WA0001.jpg", 1)
    _photo(f.photos / "Аюб" / "p2.jpg", 2)
    _photo(f.photos / "Аюб" / "p5.jpg", 5)
    _photo(f.photos / "Аюб" / "p7.jpg", 7)
    _photo(f.photos / "Сулиман С" / "p3.jpg", 3)
    _run(f)

    log = _log(f)
    assert {n: r["outcome"] for n, r in log.items()} == {
        "IMG-20260915-WA0001.jpg": "PLUS", "p2.jpg": "MINUS", "p5.jpg": "SERIAL_AMBIGUOUS",
        "p7.jpg": "REPEAT", "p3.jpg": "DIGITS_ERROR"}
    assert log["p5.jpg"]["notes"] == "номер 44444 у нескольких абонентов: A-4, A-5"
    assert log["p7.jpg"]["notes"].startswith("pre-existing entry in table")

    r = _readings(f)
    assert (r["A-1"].value, r["A-1"].date, r["A-1"].source, r["A-1"].photo) == \
        ("1200", "15 09 2026", "auto", "IMG-20260915-WA0001.jpg")             # дата из имени
    assert (r["A-2"].value, r["A-2"].date) == ("4000", TODAY)
    assert r["A-6"].value == "8" and r["A-6"].source == "table"       # не тронуто
    assert "A-4" not in r and "A-5" not in r

    res = f.results
    assert (res / "Аюб" / "plus" / "A-1.jpg").is_file()
    assert (res / "Аюб" / "question" / "serial_ambiguous" / "p5.jpg").is_file()
    assert (res / "Сулиман С" / "question" / "digits_error" / "A-3.jpg").is_file()
    assert list((res / "run_logs").glob("run_*.txt"))
    assert list(f.backups.glob("*/gmr.sqlite"))                      # копия базы перед прогоном

    with MonthDB(f.db) as db:
        changes = [(c["action"], c["account"], c["new"], c["note"]) for c in db.changes()]
    assert ("показание записано", "A-1", "1200", "IMG-20260915-WA0001.jpg") in changes

    ws = openpyxl.load_workbook(f.export_xlsx).active
    rows = {r[1].value: [c.value for c in r] for r in ws.iter_rows(min_row=2)}
    assert rows["A-1"][3] == 1200 and rows["A-2"][3] == 4000 and rows["A-6"][3] == 8


def test_rerun_month_leaves_readings(tmp_path, ocr):
    f = _month(tmp_path)
    _photo(f.photos / "Аюб" / "p1.jpg", 1)
    _run(f)
    _run(f)
    rows = [r for r in _log_list(f) if r["original_filename"] == "p1.jpg"]
    assert [r["outcome"] for r in rows] == ["PLUS", "REPEAT"]
    assert _readings(f)["A-1"].value == "1200"


def test_crash_mid_run_keeps_readings_in_month(tmp_path, ocr):
    # в режиме месяца исправлена ошибка test_known_issues::test_reading_survives_crash_and_rerun
    f = _month(tmp_path)
    _photo(f.photos / "Аюб" / "p1.jpg", 1)
    _photo(f.photos / "Аюб" / "p2.jpg", 2)
    ocr.crash_on = 2                                               # Ctrl+C на втором фото
    with pytest.raises(KeyboardInterrupt):
        _run(f)
    assert _readings(f)["A-1"].value == "1200"                     # первое уже в базе
    ocr.crash_on = None
    _run(f)
    r = _readings(f)
    assert r["A-1"].value == "1200" and r["A-2"].value == "4000"


def test_excel_open_does_not_lose_readings_in_month(tmp_path, ocr, monkeypatch):
    # в режиме месяца исправлена ошибка test_known_issues::test_readings_survive_table_locked_by_excel
    f = _month(tmp_path)
    _photo(f.photos / "Аюб" / "p1.jpg", 1)
    real_replace = month.os.replace

    def excel_open(src, dst):
        if Path(dst).name == "показания.xlsx":
            raise PermissionError(13, "файл занят", str(dst))
        return real_replace(src, dst)

    monkeypatch.setattr(month.os, "replace", excel_open)
    _run(f)                                                        # прогон не падает
    assert _readings(f)["A-1"].value == "1200"
    journal = sorted((f.results / "run_logs").glob("run_*.txt"))[-1].read_text(encoding="utf-8")
    assert "Закройте" in journal and "Показания сохранены в базе" in journal
    monkeypatch.setattr(month.os, "replace", real_replace)        # Excel закрыли
    month.export_month(str(f.root))
    rows = {r[1].value: r[3].value for r in openpyxl.load_workbook(f.export_xlsx).active.iter_rows(min_row=2)}
    assert rows["A-1"] == 1200


def test_run_requires_month(tmp_path):
    with pytest.raises(reader.NotAMonth, match="Месяц не создан"):
        reader.run_pipeline(reader.PipelineConfig(month_dir=str(tmp_path / "нет")))


def test_cli_process_month(tmp_path, ocr, capsys):
    f = _month(tmp_path)
    _photo(f.photos / "Аюб" / "p1.jpg", 1)
    assert gmr.main(["process", str(f.root)]) == 0
    assert _readings(f)["A-1"].value == "1200"
    assert gmr.main(["process", str(tmp_path / "нет")]) == 1
    assert "Месяц не создан" in capsys.readouterr().out
    with pytest.raises(SystemExit):
        gmr.main(["process", str(f.root), "--table", "t.csv"])


# ─── Номер у нескольких абонентов и в старом режиме (таблица) ───────────────

def test_ambiguous_serial_in_table_mode(base_config):
    import pandas as pd
    df = pd.DataFrame({
        base_config.col_serial: ["44444", "44444", "11111"],
        base_config.col_account_id: ["A-4", "A-5", "A-1"],
        base_config.col_last_reading: ["10", "20", "1000"],
        base_config.col_new_reading: ["", "", ""],
    })
    digit_crops, arrays = make_digit_crops_with_centers([0, 40, 80, 120, 160])
    preds = {id(a): (d, 0.95) for a, d in zip(arrays, "00500")}
    res = process_photo("x.jpg", df, base_config, FakeMeterDetector(make_meter_crops()),
                        FakeDigitDetector(digit_crops), FakeDigitOCR(preds), FakeSerialOCR("44444", 0.95))
    assert res.outcome == Outcome.SERIAL_AMBIGUOUS
    assert res.error_detail == "номер 44444 у нескольких абонентов: A-4, A-5"
    assert res.reading_str == "00500" and res.account_id is None
