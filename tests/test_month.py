"""
Месяц (этап 2.2a): загрузка таблицы компании в базу, обновление посреди
месяца, перенос старого CSV-лога, выгрузка показания.xlsx и лог.csv.

tests/data/register_sample.xls — ВЫДУМАННЫЕ абоненты; файл повторяет
особенности настоящего реестра: лицевой счёт числом и текстом, номер с
ведущими нулями, номера с лишними знаками по краям, номер у двух абонентов,
дата в ячейке-дате, пустая строка, строка без счёта, повтор счёта.
"""
import csv
import json
import sqlite3
from pathlib import Path

import openpyxl
import pytest

import gmr
from src.gmr.application import month
from src.gmr.storage import LOG_COLUMNS, load_log, save_log
from src.gmr.storage.month import MonthDB, MonthFolder
from src.gmr.storage.register import clean_serial, read_register

ROOT = Path(__file__).resolve().parent.parent
SAMPLE = ROOT / "tests" / "data" / "register_sample.xls"
HEAD = ["Лицевой счет", "Абонент", "Адрес", "ТИП", "Дата", "Последние показания",
        "Текущие показания", "Разница", "Пломба ввода", "Сейф пломбы", "Номер счетчика",
        "Номер телефона", "Абонент.Участок"]


def _xlsx(path: Path, rows: list[dict], head=HEAD):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Лист1"
    ws.append(head)
    for r in rows:
        ws.append([r.get(h) for h in head])
    wb.save(path)
    return path


def _db(folder: Path) -> MonthDB:
    return MonthDB(MonthFolder(folder).db)


# ─── Чтение таблицы ──────────────────────────────────────────────────────────

def test_read_xls_like_excel_shows_it():
    reg = read_register(str(SAMPLE))
    assert reg.columns == HEAD and reg.sheet == "Лист1"
    assert len(reg.rows) == 8                                   # пустая строка пропущена
    first = reg.rows[0]
    assert first["Лицевой счет"] == "1000000001"                 # число, а не "1000000001.0"
    assert first["Номер счетчика"] == "0012345"                  # ведущие нули
    assert first["Последние показания"] == "1000" and first["Текущие показания"] == ""
    assert reg.rows[2]["Дата"] == "20 05 2026"                   # ячейка-дата
    assert reg.rows[5]["Номер счетчика"] == "4444444"            # номер числом
    assert "Последние показания" in reg.numeric_columns
    assert "Номер телефона" not in reg.numeric_columns and "Абонент" not in reg.numeric_columns


def test_read_xlsx_and_csv_give_same_text(tmp_path):
    rows = [{"Лицевой счет": 1000000001, "Номер счетчика": "0012345", "Последние показания": 1000,
             "Текущие показания": None, "Номер телефона": "89001234567"}]
    x = read_register(str(_xlsx(tmp_path / "t.xlsx", rows)))
    assert x.rows[0]["Лицевой счет"] == "1000000001" and x.rows[0]["Номер счетчика"] == "0012345"
    assert "Номер телефона" not in x.numeric_columns
    for name, enc, sep in (("a.csv", "utf-8", ","), ("b.csv", "cp1251", ";")):
        p = tmp_path / name
        p.write_bytes(f"Лицевой счет{sep}Номер счетчика\n1000000001{sep}0012345\n".encode(enc))
        assert read_register(str(p)).rows == [{"Лицевой счет": "1000000001", "Номер счетчика": "0012345"}]


def test_read_rejects_other_formats(tmp_path):
    (tmp_path / "t.txt").write_text("x", encoding="utf-8")
    with pytest.raises(ValueError, match="xls"):
        read_register(str(tmp_path / "t.txt"))


@pytest.mark.parametrize("raw, clean", [
    ("123456.", "123456"), (",654321", "654321"), ("12345678/", "12345678"), (" 0012345 ", "0012345"),
    ("0012345", "0012345"), ("12.34", "12.34"),
])
def test_clean_serial(raw, clean):
    assert clean_serial(raw) == clean


# ─── Создание месяца ─────────────────────────────────────────────────────────

def test_create_month(tmp_path):
    rep = month.load_table(str(tmp_path / "m"), str(SAMPLE), who="тест")
    assert rep.created and rep.abonents == 6
    assert rep.serials_cleaned == [("1000000002", "123456.", "123456"), ("1000000005", ",654321", "654321")]
    assert rep.duplicate_serials == {"777777": ["1000000003", "1000000004"]}
    assert rep.no_account == 1 and rep.duplicate_accounts == ["1000000002"]
    assert rep.readings_from_table == 1
    with _db(tmp_path / "m") as db:
        ab = db.abonents()
        assert list(ab) == [f"100000000{i}" for i in range(1, 7)]             # порядок файла
        assert ab["1000000002"].serial == "123456" and ab["1000000002"].data["Абонент"] == "Пробный П.П."
        assert db.accounts_by_serial("0012345") == ["1000000001"]
        r = db.readings()["1000000002"]
        assert (r.value, r.date, r.source, r.updated_by) == ("5100", "15 05 2026", "table", "тест")
        assert [c["action"] for c in db.changes()] == ["месяц создан"]
    f = MonthFolder(tmp_path / "m")
    assert f.photos.is_dir() and f.results.is_dir()
    tables = sorted(p.name for p in f.tables.iterdir())
    assert len(tables) == 2 and tables[0].endswith("register_sample.xls") and tables[1].endswith("отчёт_загрузки.txt")
    assert SAMPLE.read_bytes() == (f.tables / tables[0]).read_bytes()        # копия, исходник не тронут


def test_missing_columns_change_nothing(tmp_path):
    bad = _xlsx(tmp_path / "bad.xlsx", [{"Лицевой счет": "1"}], head=["Лицевой счет", "Абонент"])
    with pytest.raises(ValueError, match="Номер счетчика"):
        month.load_table(str(tmp_path / "m"), str(bad))
    assert not MonthFolder(tmp_path / "m").exists()


def test_interrupted_first_load_leaves_no_half_month(tmp_path, monkeypatch):
    calls = []
    real = MonthDB.put_abonent

    def fail_on_third(self, a):
        calls.append(a)
        if len(calls) == 3:
            raise OSError("диск отключили")
        real(self, a)

    monkeypatch.setattr(MonthDB, "put_abonent", fail_on_third)
    with pytest.raises(OSError):
        month.load_table(str(tmp_path / "m"), str(SAMPLE))
    with _db(tmp_path / "m") as db:
        assert db.abonents() == {} and db.meta("created_at") is None
    assert not month.is_month(MonthFolder(tmp_path / "m"))
    monkeypatch.setattr(MonthDB, "put_abonent", real)
    assert month.load_table(str(tmp_path / "m"), str(SAMPLE)).created   # повтор — создание с нуля


# ─── Обновление таблицы посреди месяца ───────────────────────────────────────

def _updated_rows():
    reg = read_register(str(SAMPLE))
    rows = {r["Лицевой счет"]: dict(r) for r in reg.rows if r["Лицевой счет"]}
    rows.pop("1000000002")                                   # убран, с показанием
    rows.pop("1000000003")                                   # убран, без показания
    rows["1000000001"]["Номер счетчика"] = "0012346"         # новый номер
    rows["1000000005"]["Последние показания"] = "250"        # новое прошлое показание
    rows["1000000006"]["Текущие показания"] = "10"           # новое показание в таблице
    rows["1000000007"] = {"Лицевой счет": "1000000007", "Номер счетчика": "0070000",
                          "Последние показания": "700", "Абонент": "Новиков Н.Н."}
    return list(rows.values())


def test_update_table_by_account(tmp_path):
    m = str(tmp_path / "m")
    month.load_table(m, str(SAMPLE))
    rep = month.load_table(m, str(_xlsx(tmp_path / "upd.xlsx", _updated_rows())), who="тест")
    assert not rep.created
    assert rep.added == ["1000000007"]
    assert sorted(rep.removed) == ["1000000002", "1000000003"] and rep.removed_with_reading == ["1000000002"]
    assert rep.serial_changed == [("1000000001", "0012345", "0012346")] and rep.last_changed == 1
    assert rep.readings_from_table == 1
    with _db(tmp_path / "m") as db:
        ab = db.abonents()
        assert not ab["1000000002"].in_table and not ab["1000000003"].in_table   # не удалены
        assert ab["1000000007"].in_table and ab["1000000005"].last_reading == "250"
        assert db.readings()["1000000002"].value == "5100"                       # показание сохранено
        assert db.readings()["1000000006"].value == "10"
        actions = [(c["action"], c["account"], c["field"], c["old"], c["new"]) for c in db.changes()]
    assert ("изменён в таблице", "1000000001", "Номер счетчика", "0012345", "0012346") in actions
    assert ("нет в новой таблице", "1000000002", "", "", "") in actions
    assert ("абонент добавлен", "1000000007", "", "", "") in actions
    assert actions[-1][0] == "таблица обновлена"
    # база перед обновлением сохранена
    copies = list((MonthFolder(tmp_path / "m").backups).glob("*/gmr.sqlite"))
    assert len(copies) == 1
    with sqlite3.connect(copies[0]) as old:
        assert old.execute("SELECT COUNT(*) FROM abonents WHERE in_table=1").fetchone()[0] == 6


def test_update_keeps_reading_from_db_and_returns_abonent(tmp_path):
    m = str(tmp_path / "m")
    month.load_table(m, str(SAMPLE))
    rows = _updated_rows()
    month.load_table(m, str(_xlsx(tmp_path / "u1.xlsx", rows)))
    rows.append({"Лицевой счет": "1000000002", "Номер счетчика": "123456", "Последние показания": "5000",
                 "Текущие показания": "9999"})                              # вернулся, другое показание
    rep = month.load_table(m, str(_xlsx(tmp_path / "u2.xlsx", rows)))
    assert rep.returned == ["1000000002"]
    assert rep.readings_kept == [("1000000002", "5100", "9999")]
    with _db(tmp_path / "m") as db:
        assert db.readings()["1000000002"].value == "5100"


def test_same_table_twice_changes_nothing(tmp_path):
    m = str(tmp_path / "m")
    month.load_table(m, str(SAMPLE))
    rep = month.load_table(m, str(SAMPLE))
    assert (rep.added, rep.removed, rep.serial_changed, rep.last_changed, rep.readings_from_table) == \
        ([], [], [], 0, 0)


def test_account_with_dot_zero_from_pandas_csv(tmp_path):
    t = tmp_path / "t.csv"
    t.write_text("Лицевой счет,Номер счетчика,Последние показания,Текущие показания\n"
                 "1000000001.0,0012345,10,\n", encoding="utf-8")
    month.load_table(str(tmp_path / "m"), str(t))
    with _db(tmp_path / "m") as db:
        assert list(db.abonents()) == ["1000000001"]


# ─── Перенос старых данных (CSV-таблица + CSV-лог) ───────────────────────────

def test_old_csv_table_and_log_moved_into_month(tmp_path):
    table = tmp_path / "meters_table.csv"
    with open(table, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, quoting=csv.QUOTE_ALL)
        w.writerow(["Номер счетчика", "Лицевой счет", "Последние показания", "Текущие показания"])
        w.writerow(["0012345", "A-1", "1000", "1200"])
        w.writerow(["7654321", "A-2", "5000", ""])
    rows = []
    for name, outcome, verified in (("p1.jpg", "PLUS", "Оператор"), ("p2.jpg", "NO_METER", "")):
        r = {c: "" for c in LOG_COLUMNS}
        r.update(original_filename=name, outcome=outcome, source="auto", verified_by=verified,
                 account_id="A-1" if outcome == "PLUS" else "", reading="1200" if outcome == "PLUS" else "",
                 photo_hash=f"h-{name}", notes="Готово, «кавычки», запятые")
        rows.append(r)
    save_log(str(tmp_path / "meters_table_log.csv"), rows)

    rep = month.load_table(str(tmp_path / "m"), str(table))
    assert rep.log_rows == 2 and rep.readings_from_table == 1
    with _db(tmp_path / "m") as db:
        assert db.log_rows() == load_log(str(tmp_path / "meters_table_log.csv"))
        assert db.readings()["A-1"].value == "1200"
    # повторная загрузка (обновление) лог второй раз не переносит
    assert month.load_table(str(tmp_path / "m"), str(table)).log_rows == 0


# ─── Выгрузка ────────────────────────────────────────────────────────────────

def _sheet(folder):
    return openpyxl.load_workbook(MonthFolder(folder).export_xlsx).active


def test_export_xlsx_same_columns_and_types(tmp_path):
    m = tmp_path / "m"
    month.load_table(str(m), str(SAMPLE))
    rep = month.export_month(str(m))
    assert (rep.rows, rep.with_reading) == (6, 1)
    ws = _sheet(m)
    assert ws.title == "Лист1" and [c.value for c in ws[1]] == HEAD
    rows = [{h: c.value for h, c in zip(HEAD, r)} for r in ws.iter_rows(min_row=2)]
    assert [r["Лицевой счет"] for r in rows] == [f"100000000{i}" for i in range(1, 7)]   # текст
    assert rows[0]["Номер счетчика"] == "0012345" and rows[5]["Номер счетчика"] == "4444444"
    # выгрузка — таблица компании с нашими показаниями: её номер как был (очищенный —
    # только в базе, для поиска фото)
    assert rows[1]["Номер счетчика"] == "123456."
    assert rows[1]["Текущие показания"] == 5100 and rows[0]["Текущие показания"] is None
    assert rows[0]["Последние показания"] == 1000 and rows[0]["Пломба ввода"] == 11111111
    assert rows[1]["Номер телефона"] == "+7(900) 000-00-00" and rows[1]["Дата"] == "15 05 2026"
    assert rows[2]["Номер телефона"] == "89000000000"                  # текстом, как в исходнике
    assert rows[0]["Разница"] == "=G2-F2" and rows[5]["Разница"] == "=G7-F7"          # формула, как в исходнике


def test_export_skips_removed_abonents_and_uses_db_reading(tmp_path):
    m = tmp_path / "m"
    month.load_table(str(m), str(SAMPLE))
    month.load_table(str(m), str(_xlsx(tmp_path / "upd.xlsx", _updated_rows())))
    with _db(m) as db, db.transaction():
        r = db.readings()["1000000006"]
        r.value, r.date, r.source = "11", "02 10 2026", "auto"         # так будет писать reader (2.2b)
        db.put_reading(r)
    month.export_month(str(m))
    rows = {r[0].value: r for r in _sheet(m).iter_rows(min_row=2)}
    assert set(rows) == {"1000000001", "1000000004", "1000000005", "1000000006", "1000000007"}
    six = {h: c.value for h, c in zip(HEAD, rows["1000000006"])}
    assert six["Текущие показания"] == 11 and six["Дата"] == "02 10 2026"


def test_export_log_csv_like_old_log(tmp_path):
    m = tmp_path / "m"
    month.load_table(str(m), str(SAMPLE))
    row = {c: "" for c in LOG_COLUMNS}
    row.update(original_filename="p.jpg", outcome="PLUS", notes="а, «б»")
    with _db(m) as db, db.transaction():
        db.append_log_rows([row])
    month.export_month(str(m))
    assert load_log(str(MonthFolder(m).export_log)) == [row]


def test_export_when_file_open_in_excel(tmp_path, monkeypatch):
    m = tmp_path / "m"
    month.load_table(str(m), str(SAMPLE))
    month.export_month(str(m))
    before = MonthFolder(m).export_xlsx.read_bytes()

    def locked(src, dst):
        raise PermissionError(13, "файл занят", str(dst))

    monkeypatch.setattr(month.os, "replace", locked)
    with pytest.raises(month.ExportLocked, match="Закройте"):
        month.export_month(str(m))
    assert MonthFolder(m).export_xlsx.read_bytes() == before                 # старая выгрузка цела
    assert not list(m.glob(".*tmp*"))                                         # временный файл убран


def test_export_without_month(tmp_path):
    with pytest.raises(ValueError, match="Месяц не создан"):
        month.export_month(str(tmp_path / "нет"))


# ─── Команды ─────────────────────────────────────────────────────────────────

def test_cli_month_and_export(tmp_path, capsys):
    m = str(tmp_path / "Октябрь_2026")
    assert gmr.main(["month", m]) == 0
    assert "Месяц не создан" in capsys.readouterr().out
    assert gmr.main(["month", m, "--table", str(SAMPLE)]) == 0
    out = capsys.readouterr().out
    assert "МЕСЯЦ СОЗДАН" in out and "показания.xlsx" in out
    assert gmr.main(["month", m]) == 0
    out = capsys.readouterr().out
    assert "абонентов в таблице: 6" in out and "с показанием: 1 — программа 0, оператор 0, было в таблице 1" in out
    assert gmr.main(["export", m]) == 0
    assert gmr.main(["export", str(tmp_path / "нет")]) == 1
    assert "ОШИБКА" in capsys.readouterr().out


def test_cli_month_bad_table_is_error(tmp_path, capsys):
    bad = _xlsx(tmp_path / "bad.xlsx", [], head=["Лицевой счет"])
    assert gmr.main(["month", str(tmp_path / "m"), "--table", str(bad)]) == 1
    assert "нет столбцов" in capsys.readouterr().out


def test_numeric_columns_stored(tmp_path):
    month.load_table(str(tmp_path / "m"), str(SAMPLE))
    with _db(tmp_path / "m") as db:
        assert "Номер телефона" not in json.loads(db.meta("numeric_columns"))


def test_failed_transaction_leaves_nothing_even_if_db_used_again(tmp_path):
    with MonthDB(tmp_path / "gmr.sqlite") as db:
        with pytest.raises(RuntimeError):
            with db.transaction():
                db.set_meta("половина", "1")
                raise RuntimeError("сбой посреди записи")
        with db.transaction():
            db.set_meta("следующая", "2")
        assert db.meta("половина") is None and db.meta("следующая") == "2"


# ─── Отчёт загрузки: где лицевой счёт, где номер; полные списки — в файле ────

def test_report_labels_account_and_serial(tmp_path):
    text = month.load_table(str(tmp_path / "m"), str(SAMPLE)).text()
    assert "л/с 1000000002: номер счётчика «123456.» → «123456»" in text
    assert "номер счётчика 777777 — у л/с 1000000003, 1000000004" in text


def test_long_lists_full_in_report_file(tmp_path):
    rows = [{"Лицевой счет": f"20000000{i:02d}", "Номер счетчика": f"5000{i:02d}.",
             "Последние показания": "1"} for i in range(25)]
    rep = month.load_table(str(tmp_path / "m"), str(_xlsx(tmp_path / "t.xlsx", rows)))
    screen = rep.text()
    assert screen.count("номер счётчика «5000") == 20
    assert f"… и ещё 5 — полный список в файле {rep.report_file}" in screen
    saved = Path(rep.report_file).read_text(encoding="utf-8")
    assert saved.count("номер счётчика «5000") == 25 and "… и ещё" not in saved
