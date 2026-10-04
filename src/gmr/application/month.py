"""
src/gmr/application/month.py — месяц: загрузка таблицы компании в базу,
обновление посреди месяца, выгрузка показания.xlsx (этап 2.2a, 2026-10-01).

Решения владельца 2026-10-01 (docs/BACKLOG.md, этап 2):
  - одна папка — один месяц (src/gmr/storage/month.py);
  - исходный файл компании не меняется; результат — показания.xlsx, те же
    столбцы в том же порядке;
  - номера счётчиков с лишними знаками по краям очищаются при загрузке;
  - показания, уже заполненные в таблице при загрузке, считаются внесёнными
    до программы;
  - обновлённая таблица сверяется по лицевому счёту: новые абоненты
    добавляются, исчезнувшие помечаются «нет в таблице» (не удаляются),
    у остальных обновляются номер и прошлое показание; показание из базы
    важнее показания в новой таблице;
  - при создании месяца старый CSV-лог (`<таблица>_log.csv`) переносится в
    базу целиком.

Пока (2.2a) reader.py и program2.py работают по-старому, с CSV; на базу они
переходят на этапах 2.2b и 2.3.
"""
import csv
import getpass
import json
import os
import re
import shutil
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Optional

from src.gmr.domain.config import PipelineConfig
from src.gmr.domain.models import QUESTION_REASONS
from src.gmr.storage.backup import backup_sqlite
from src.gmr.storage.log_store import LOG_COLUMNS, load_log, log_path_for
from src.gmr.storage.month import Abonent, MonthDB, MonthFolder, Reading, now_text
from src.gmr.storage.register import clean_serial, read_register

_LIST_LIMIT = 20          # сколько строк списка показывать в отчёте


# журнал изменений: оператор исправил номер счётчика в базе («Серийник в базе с
# ошибкой», 2026-10-01); при обновлении таблицы такое исправление сохраняется
SERIAL_FIX_ACTION = "номер исправлен оператором"


class ExportLocked(Exception):
    """Файл выгрузки открыт в Excel — Windows не даёт его заменить."""


class NotAMonth(ValueError):
    """Папка — не месяц: база не создана (сначала `gmr.py month … --table …`)."""


def normalize_account(text: str) -> str:
    """Лицевой счёт как текст: без пробелов и без «.0» (так его пишет pandas)."""
    text = (text or "").strip()
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".")[0]
    return text


def _who() -> str:
    try:
        return getpass.getuser()
    except Exception:
        return "?"


def _lines(title: str, items: list, fmt, limit: Optional[int], full: str) -> list[str]:
    if not items:
        return []
    out = [f"{title}: {len(items)}"]
    shown = items if limit is None else items[:limit]
    out += [f"    {fmt(x)}" for x in shown]
    if len(items) > len(shown):
        out.append(f"    … и ещё {len(items) - len(shown)} — полный список в файле {full}")
    return out


# ─── Загрузка таблицы ────────────────────────────────────────────────────────

@dataclass
class LoadReport:
    folder: str
    table: str
    created: bool
    abonents: int = 0
    readings_from_table: int = 0
    added: list = field(default_factory=list)
    removed: list = field(default_factory=list)
    removed_with_reading: list = field(default_factory=list)
    returned: list = field(default_factory=list)
    serial_changed: list = field(default_factory=list)        # (счёт, было, стало)
    last_changed: int = 0
    serials_cleaned: list = field(default_factory=list)       # (счёт, в таблице, очищен)
    duplicate_serials: dict = field(default_factory=dict)     # номер -> [счета]
    no_account: int = 0
    duplicate_accounts: list = field(default_factory=list)
    readings_kept: list = field(default_factory=list)         # (счёт, в базе, в таблице)
    serial_fixes_kept: list = field(default_factory=list)     # (счёт, в таблице, исправлен оператором)
    log_rows: int = 0

    report_file: str = ""                                      # полный отчёт в таблицы/

    def text(self, limit: Optional[int] = _LIST_LIMIT) -> str:
        """Отчёт загрузки. limit — сколько строк каждого списка показать
        (None — все, так отчёт пишется в файл)."""
        def lines(title, items, fmt):
            return _lines(title, items, fmt, limit, self.report_file)

        out = [
            "=" * 60,
            ("МЕСЯЦ СОЗДАН" if self.created else "ТАБЛИЦА ОБНОВЛЕНА") + f": {self.folder}",
            f"Таблица: {self.table}",
            "=" * 60,
            f"Абонентов в таблице: {self.abonents}",
        ]
        if self.readings_from_table:
            out.append(f"Показаний уже было в таблице (считаются внесёнными до программы): "
                       f"{self.readings_from_table}")
        if not self.created:
            out.append(f"Новых абонентов: {len(self.added)}")
            out += lines("Нет в новой таблице (в выгрузку не попадут)", self.removed, lambda x: f"л/с {x}")
            out += lines("  из них уже с показанием этого месяца", self.removed_with_reading,
                         lambda x: f"л/с {x}")
            out += lines("Вернулись в таблицу", self.returned, lambda x: f"л/с {x}")
            out += lines("Изменился номер счётчика", self.serial_changed,
                         lambda x: f"л/с {x[0]}: номер счётчика «{x[1]}» → «{x[2]}»")
            if self.last_changed:
                out.append(f"Изменилось прошлое показание: {self.last_changed}")
            out += lines("Номер, исправленный оператором, сохранён (в новой таблице прежний)",
                         self.serial_fixes_kept,
                         lambda x: f"л/с {x[0]}: в таблице «{x[1]}», в базе «{x[2]}»")
            out += lines("В новой таблице другое показание — оставлено показание из базы",
                         self.readings_kept, lambda x: f"л/с {x[0]}: в базе {x[1]}, в таблице {x[2]}")
        out += lines("Номер счётчика с лишними знаками по краям — очищен в базе (исходный файл не менялся)",
                     self.serials_cleaned, lambda x: f"л/с {x[0]}: номер счётчика «{x[1]}» → «{x[2]}»")
        out += lines("Один номер счётчика у нескольких абонентов", list(self.duplicate_serials.items()),
                     lambda x: f"номер счётчика {x[0]} — у л/с {', '.join(x[1])}")
        if self.no_account:
            out.append(f"Строк без лицевого счёта (пропущены): {self.no_account}")
        out += lines("Лицевой счёт повторяется (взята первая строка)", self.duplicate_accounts,
                     lambda x: f"л/с {x}")
        if self.log_rows:
            out.append(f"Перенесён старый лог обработки: {self.log_rows} строк")
        if limit is not None and self.report_file:
            out.append(f"Полный отчёт: {self.report_file}")
        out.append("=" * 60)
        return "\n".join(out)


def load_table(month_dir: str, table_path: str, config: Optional[PipelineConfig] = None,
               who: Optional[str] = None, import_old_log: bool = True) -> LoadReport:
    """Создаёт месяц из таблицы компании или загружает её обновлённую версию.
    Ошибка ValueError — если таблица не читается или в ней нет нужных столбцов;
    тогда база не меняется. import_old_log=False — не переносить лог старого
    режима (<таблица>_log.csv рядом с таблицей) при создании месяца: окно
    оператора спрашивает об этом (этап 4)."""
    cfg = config or PipelineConfig()
    who = who or _who()
    folder = MonthFolder(Path(month_dir))
    reg = read_register(table_path)
    required = [cfg.col_account_id, cfg.col_serial, cfg.col_last_reading, cfg.col_new_reading]
    missing = [c for c in required if c not in reg.columns]
    if missing:
        raise ValueError(
            "В таблице нет столбцов: " + ", ".join(f"«{c}»" for c in missing)
            + ". Есть: " + ", ".join(f"«{c}»" for c in reg.columns)
        )

    rep = LoadReport(folder=str(folder.root), table=str(table_path), created=False)

    # строки таблицы → абоненты (в порядке файла)
    incoming: dict[str, Abonent] = {}
    for order, row in enumerate(reg.rows):
        account = normalize_account(row.get(cfg.col_account_id, ""))
        if not account:
            rep.no_account += 1
            continue
        if account in incoming:
            rep.duplicate_accounts.append(account)
            continue
        raw = row.get(cfg.col_serial, "")
        serial = clean_serial(raw)
        if serial != raw:
            rep.serials_cleaned.append((account, raw, serial))
        data = dict(row)
        data[cfg.col_account_id] = account
        incoming[account] = Abonent(account, serial, row.get(cfg.col_last_reading, ""), True, order, data)

    folder.root.mkdir(parents=True, exist_ok=True)
    with MonthDB(folder.db) as db:
        # месяц создан, только если создание дошло до конца (запись created_at
        # делается в той же транзакции, что и абоненты)
        created = rep.created = db.meta("created_at") is None
        if not created:
            backup_sqlite(folder.db, folder.backups)
        with db.transaction():
            old = db.abonents()
            readings = db.readings()
            stamp = now_text()
            fixes = {c["account"]: c["new"] for c in db.changes() if c["action"] == SERIAL_FIX_ACTION}
            for account, a in incoming.items():
                fixed = fixes.get(account)
                if fixed is not None and a.serial != fixed:
                    rep.serial_fixes_kept.append((account, a.serial, fixed))
                    a.serial = fixed
                prev = old.get(account)
                if prev is None:
                    if not created:
                        rep.added.append(account)
                        db.add_change(who, "абонент добавлен", account)
                else:
                    if not prev.in_table:
                        rep.returned.append(account)
                        db.add_change(who, "абонент вернулся в таблицу", account)
                    if prev.serial != a.serial:
                        rep.serial_changed.append((account, prev.serial, a.serial))
                        db.add_change(who, "изменён в таблице", account, cfg.col_serial, prev.serial, a.serial)
                    if prev.last_reading != a.last_reading:
                        rep.last_changed += 1
                        db.add_change(who, "изменён в таблице", account, cfg.col_last_reading,
                                      prev.last_reading, a.last_reading)
                db.put_abonent(a)

                in_table = a.data.get(cfg.col_new_reading, "").strip()
                if in_table:
                    have = readings.get(account)
                    if have is None:
                        db.put_reading(Reading(account, in_table, a.data.get(cfg.col_date, ""),
                                               "table", "", stamp, who))
                        rep.readings_from_table += 1
                    elif have.value != in_table:
                        rep.readings_kept.append((account, have.value, in_table))

            for account, prev in old.items():
                if account not in incoming and prev.in_table:
                    prev.in_table = False
                    db.put_abonent(prev)
                    rep.removed.append(account)
                    if account in readings:
                        rep.removed_with_reading.append(account)
                    db.add_change(who, "нет в новой таблице", account)

            db.set_meta("columns", json.dumps(reg.columns, ensure_ascii=False))
            db.set_meta("sheet", reg.sheet)
            db.set_meta("numeric_columns", json.dumps(reg.numeric_columns, ensure_ascii=False))
            db.set_meta("table", Path(table_path).name)
            db.set_meta("table_loaded_at", stamp)
            if created:
                db.set_meta("created_at", stamp)
                old_log = Path(log_path_for(table_path))
                if import_old_log and old_log.is_file():
                    rows = load_log(str(old_log))
                    db.append_log_rows(rows)
                    rep.log_rows = len(rows)
            db.add_change(who, "месяц создан" if created else "таблица обновлена",
                          note=Path(table_path).name)

        by_serial: dict[str, list[str]] = {}
        for a in db.abonents(only_in_table=True).values():
            if a.serial:
                by_serial.setdefault(a.serial, []).append(a.account)
        rep.duplicate_serials = {s: accs for s, accs in by_serial.items() if len(accs) > 1}
        rep.abonents = sum(1 for _ in db.abonents(only_in_table=True))

    for d in (folder.photos, folder.results, folder.tables):
        d.mkdir(exist_ok=True)
    stamp_file = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    shutil.copy2(table_path, folder.tables / f"{stamp_file}_{Path(table_path).name}")
    report = folder.tables / f"{stamp_file}_отчёт_загрузки.txt"
    rep.report_file = str(report)
    report.write_text(rep.text(limit=None) + "\n", encoding="utf-8")       # все списки целиком
    return rep


# ─── Выгрузка ────────────────────────────────────────────────────────────────

@dataclass
class ExportReport:
    path: str
    rows: int
    with_reading: int
    log_path: str
    log_rows: int

    def text(self) -> str:
        return (f"Выгрузка: {self.path}\n"
                f"  абонентов: {self.rows}, с показанием: {self.with_reading}, "
                f"без показания: {self.rows - self.with_reading}\n"
                f"Лог: {self.log_path} ({self.log_rows} строк)")


_INT = re.compile(r"-?(0|[1-9]\d*)")
_FLOAT = re.compile(r"-?(0|[1-9]\d*)\.\d+")


def _cell_value(text: str, as_text: bool):
    if as_text or not text:
        return text
    if _INT.fullmatch(text):
        return int(text)
    if _FLOAT.fullmatch(text):
        return float(text)
    return text


def _replace(tmp: Path, target: Path) -> None:
    try:
        os.replace(tmp, target)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise ExportLocked(
            f"Файл открыт в Excel: {target}. Закройте его и выгрузите заново: "
            f"python gmr.py export \"{target.parent}\""
        ) from None


def export_month(month_dir: str, config: Optional[PipelineConfig] = None) -> ExportReport:
    """показания.xlsx (столбцы таблицы компании, абоненты «в таблице», показания
    из базы, «Разница» — формулой, как в исходнике) и лог.csv. Файлы
    заменяются целиком; если показания.xlsx открыт — ExportLocked, база не
    меняется."""
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    cfg = config or PipelineConfig()
    folder = MonthFolder(Path(month_dir))
    if not is_month(folder):
        raise ValueError(f"Месяц не создан: {folder.root}. Сначала: "
                         f"python gmr.py month \"{folder.root}\" --table <таблица компании>")
    with MonthDB(folder.db) as db:
        columns = db.columns()
        sheet = db.meta("sheet") or "Лист1"
        numeric = set(json.loads(db.meta("numeric_columns") or "[]"))
        abonents = list(db.abonents(only_in_table=True).values())
        readings = db.readings()
        log_rows = db.log_rows()

    letter = {c: get_column_letter(i) for i, c in enumerate(columns, start=1)}
    # лицевой счёт и номер — всегда текст (ведущие нули); остальное — как было
    # в исходнике; показания — числом
    text_cols = ({cfg.col_account_id, cfg.col_serial} | (set(columns) - numeric)) - {cfg.col_new_reading}
    formula = cfg.col_difference in letter and cfg.col_new_reading in letter and cfg.col_last_reading in letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    widths = {c: len(c) for c in columns}
    with_reading = 0
    for n, a in enumerate(abonents, start=2):
        r = readings.get(a.account)
        with_reading += r is not None
        row = []
        for col in columns:
            text = a.data.get(col, "")
            if col == cfg.col_new_reading:
                text = r.value if r else ""
            elif col == cfg.col_date and r is not None and r.date:
                text = r.date
            if col == cfg.col_difference and formula:
                row.append(f"={letter[cfg.col_new_reading]}{n}-{letter[cfg.col_last_reading]}{n}")
                continue
            widths[col] = max(widths[col], len(text))
            row.append(_cell_value(text, col in text_cols))
        ws.append(row)
    for col in columns:
        ws.column_dimensions[letter[col]].width = min(widths[col] + 2, 50)
        if col in text_cols:
            for (cell,) in ws.iter_rows(min_row=2, min_col=columns.index(col) + 1,
                                        max_col=columns.index(col) + 1):
                cell.number_format = "@"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    tmp = folder.root / ".показания.tmp.xlsx"
    wb.save(str(tmp))
    _replace(tmp, folder.export_xlsx)

    tmp = folder.root / ".лог.tmp.csv"
    with open(tmp, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=LOG_COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(log_rows)
    _replace(tmp, folder.export_log)

    return ExportReport(str(folder.export_xlsx), len(abonents), with_reading,
                        str(folder.export_log), len(log_rows))


# ─── Сводка ──────────────────────────────────────────────────────────────────

def is_month(folder: MonthFolder) -> bool:
    """Месяц создан: база есть, и загрузка таблицы в неё дошла до конца."""
    if not folder.exists():
        return False
    with MonthDB(folder.db) as db:
        return db.meta("created_at") is not None


WITHOUT_READING_LIST = "без_показаний.csv"     # в папке месяца — пишет month_summary
_SOURCE_NAMES = (("auto", "программа"), ("manual", "оператор"), ("table", "было в таблице"))
_PHOTO_EXTS = (".jpg", ".jpeg", ".png")


def waiting_photos(results: Path) -> dict[str, int]:
    """Сколько фото ждут оператора: файлы в результат/<контролёр>/question/<причина>/,
    по причинам в порядке очереди (только причины, где что-то есть)."""
    counts = {}
    for reason in QUESTION_REASONS:
        n = sum(1 for p in Path(results).glob(f"*/question/{reason}/*")
                if p.is_file() and p.suffix.lower() in _PHOTO_EXTS)
        if n:
            counts[reason] = n
    return counts


def _write_without_reading(path: Path, columns: list[str], abonents: list) -> Optional[str]:
    """Список абонентов без показания (все столбцы таблицы компании; «;» — для
    Excel). Возвращает текст ошибки, если файл не записался (открыт в Excel)."""
    try:
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(columns)
            for a in abonents:
                w.writerow([a.data.get(c, "") for c in columns])
    except OSError as e:
        return f"не записан ({e}) — если открыт в Excel, закройте"
    return None


def month_summary(month_dir: str) -> str:
    """Итог месяца (этап 3, пункт 6а, решение владельца 2026-10-03): сколько
    показаний записано и кем, кто без показания (полный список — в
    без_показаний.csv), сколько фото ждут оператора и по каким причинам,
    сколько номеров исправил оператор."""
    folder = MonthFolder(Path(month_dir))
    if not is_month(folder):
        return (f"Месяц не создан: {folder.root}\n"
                f"Создать: python gmr.py month \"{folder.root}\" --table <таблица компании>")
    with MonthDB(folder.db) as db:
        in_table = db.abonents(only_in_table=True)
        everyone = db.abonents()
        readings = db.readings()
        columns = db.columns()
        log_rows = len(db.log_rows())
        serial_fixes = {c["account"] for c in db.changes() if c["action"] == SERIAL_FIX_ACTION}
        meta = (db.meta("created_at"), db.meta("table"), db.meta("table_loaded_at"))
    with_reading = [a for a in in_table if a in readings]
    by_source = {src: sum(1 for a in with_reading if readings[a].source == src) for src, _ in _SOURCE_NAMES}
    without = [ab for a, ab in in_table.items() if a not in readings]
    list_path = folder.root / WITHOUT_READING_LIST
    list_error = _write_without_reading(list_path, columns, without)
    waiting = waiting_photos(folder.results)
    lines = [
        f"Месяц: {folder.root}",
        f"  создан: {meta[0]}, таблица: {meta[1]} (загружена {meta[2]})",
        f"  абонентов в таблице: {len(in_table)}",
        f"  с показанием: {len(with_reading)} — "
        + ", ".join(f"{name} {by_source[src]}" for src, name in _SOURCE_NAMES),
        f"  без показания: {len(without)} — список: {list_path}" + (f" ({list_error})" if list_error else ""),
        f"  нет в таблице (после обновления): {len(everyone) - len(in_table)}",
        f"  ждут оператора: {sum(waiting.values())} фото"
        + (" — " + ", ".join(f"{QUESTION_REASONS[r]} {n}" for r, n in waiting.items()) if waiting else ""),
        f"  номера счётчиков исправлены оператором: {len(serial_fixes)}"
        + (" (для компании — db_serial_fix.csv)" if serial_fixes else ""),
        f"  строк в логе обработки: {log_rows}",
    ]
    return "\n".join(lines)
